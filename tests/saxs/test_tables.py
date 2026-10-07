"""Tests for saxs.core.tables: transliteration, the three writers, the model-free sheet and the fit-workflow sheets."""
from __future__ import annotations

import numpy as np
import pandas as pd
import pytest

from saxs.core import tables as tb


def test_ascii_transliteration_covers_the_characters_excel_garbles():
    assert tb.ascii_transliterate("0.15–0.20 Å⁻¹ ± 2 × 10⁻³ µm, R g, d ≈ 4, T≥3 °C") == \
        "0.15-0.20 A^-1 +- 2 x 10^-3 mum, R g, d ~ 4, T>=3  degC"
    assert tb.ascii_transliterate("Bi₂O₃ (mol%)") == "Bi2O3 (mol%)" and tb.ascii_transliterate("cm²·s⁻¹") == "cm^2.s^-1"
    assert tb.ascii_transliterate("naïve café") == "naive cafe"
    assert tb.ascii_transliterate(None) == "" and tb.ascii_transliterate(3.5) == "3.5"
    out = tb.ascii_transliterate("日本語 — test")
    assert out.isascii() and "test" in out


def test_format_cell_and_markdown():
    df = pd.DataFrame({"Sample": ["a", "b"], "Rg (nm)": [12.3456789, np.nan], "ok": [True, False], "note": [None, "x"]})
    md = tb.dataframe_to_markdown(df, notes=["first note", "second"])
    lines = md.splitlines()
    assert lines[0] == "| Sample | Rg (nm) | ok | note |" and lines[1] == "|---|---|---|---|"
    assert lines[2] == "| a | 12.3457 | yes | - |" and lines[3] == "| b | - | no | x |"
    assert "- first note" in md and "- second" in md
    assert tb.format_cell(1e-7) == "1e-07" and tb.format_cell(np.float64(2.5)) == "2.5" and tb.format_cell(pd.NA) == "-"


def test_write_tables_xlsx_csv_md(tmp_path):
    fits = pd.DataFrame({"Sample": ["s1", "s2"], "Rg₂ (nm)": [12.0, None], "Loading": ["opaque", "marginal"], "note": ["≈ 4 – 5", None]})
    mf = pd.DataFrame({"Sample": ["s1", "s2"], "plateau 0.15–0.20": [1e-3, 2e-3]})
    base = str(tmp_path / "out" / "tables")
    paths = tb.write_tables({"fits": fits, "model-free": mf}, base, notes=["a note with an en-dash –"])
    assert set(paths) == {"xlsx", "csv", "md"}
    import openpyxl
    wb = openpyxl.load_workbook(paths["xlsx"])
    assert wb.sheetnames == ["fits", "model-free", "notes"]
    ws = wb["fits"]
    assert ws.freeze_panes == "B2"
    assert ws["A1"].value == "Sample" and ws["B1"].value == "Rg₂ (nm)"       # xlsx keeps the unicode
    assert 8 <= ws.column_dimensions["A"].width <= 30
    assert wb["notes"]["A2"].value.startswith("a note")
    raw = open(paths["csv"], "rb").read()
    assert raw.startswith(b"\xef\xbb\xbf")                                    # BOM
    text = raw.decode("utf-8-sig")
    assert text.isascii() and "Rg2 (nm)" in text and "~ 4 - 5" in text
    md = open(paths["md"], encoding="utf-8").read()
    assert md.startswith("| Sample | Rg₂ (nm) |") and "| s2 | - | marginal | - |" in md and "- a note" in md
    only_md = tb.write_tables({"x": fits}, str(tmp_path / "x"), formats=("md",), md_sheet="x")
    assert set(only_md) == {"md"}
    with pytest.raises(ValueError):
        tb.write_tables({}, str(tmp_path / "none"))


def test_modelfree_sheet():
    d = pd.DataFrame({"sample": ["a"], "cls": ["opaque"], "n_low": [4.734], "n_at_0.02": [-4.1], "n_at_0.04": [-3.94], "C_hi": [0.0123456],
                      "ratio_hi": [12.34], "K4_over_rhoe2": [1.23e-6], "shoulder_size_nm": [np.nan], "has_maximum": [False]})
    mf = tb.modelfree_sheet(d, extra={"x": {"a": 5}})
    r = mf.iloc[0]
    assert r["Sample"] == "a" and r["Low-q exponent"] == 4.73 and r["Local exponent at 0.04"] == 3.94 and r["Local exponent at 0.02"] == 4.1
    assert r["High-q plateau"] == 0.0123 and r["Sample/(T x empty) at high q"] == 12.3 and r["Porod shoulder (nm)"] is None
    assert r["Maximum in I"] == "no" and r["x"] == 5
    assert list(mf.columns).index("Local exponent at 0.02") < list(mf.columns).index("Local exponent at 0.04")


def test_workflow_sheets_and_writer(tmp_path):
    import fit_synthetic as fsyn
    res = fsyn.small_result()
    desc = pd.DataFrame({"sample": list(res.curves), "cls": "", "n_low": 4.2})
    sheets = tb.workflow_sheets(res, extra={"B_struct": {"x": 2.0}}, descriptors=desc)
    assert list(sheets) == ["summary", "models", "parameters", "model-free"]
    assert sheets["summary"].columns[0] == "x" and len(sheets["summary"]) == len(res.curves)
    paths = tb.write_workflow_tables(res, str(tmp_path / "w"), descriptors=desc)
    assert {"xlsx", "csv", "md", "csv:models", "csv:parameters", "csv:model-free"} <= set(paths)
    import openpyxl
    assert openpyxl.load_workbook(paths["xlsx"]).sheetnames == ["summary", "models", "parameters", "model-free", "notes"]
    assert paths["csv:model-free"].endswith("w_model_free.csv")
